import io
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from typing import List, Tuple

from pypdf import PdfReader
from docx import Document
from PIL import Image

from app import s3


IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff", "tif"}
TEXT_EXTENSIONS = {
    "txt", "md", "markdown", "csv", "json", "xml", "yaml", "yml",
    "py", "js", "ts", "tsx", "jsx", "java", "cpp", "c", "h", "hpp",
    "cs", "go", "rs", "rb", "php", "html", "css", "sql", "sh",
}


@dataclass
class ExtractedContent:
    text: str
    images_base64: List[Tuple[str, str]]  # (mime_type, base64)
    notes: List[str]


def _read_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "cp1251", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _resize_image(data: bytes, max_width: int = 1280) -> bytes:
    try:
        img = Image.open(io.BytesIO(data))
        img = img.convert("RGB") if img.mode in ("RGBA", "P") else img
        if img.width > max_width:
            ratio = max_width / img.width
            new_height = int(img.height * ratio)
            img = img.resize((max_width, new_height), Image.Resampling.LANCZOS)
        buffer = io.BytesIO()
        img.save(buffer, format="PNG")
        return buffer.getvalue()
    except Exception:
        return data


def _image_to_base64(data: bytes, mime: str = "image/png") -> str:
    import base64
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


def _is_odt(content: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if "mimetype" in zf.namelist():
                mimetype = zf.read("mimetype").decode("utf-8", errors="ignore").strip()
                return "opendocument" in mimetype
    except Exception:
        pass
    return False


def _extract_odt(content: bytes) -> ExtractedContent:
    text_parts: List[str] = []
    images: List[Tuple[str, str]] = []
    notes: List[str] = []

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            if "content.xml" in zf.namelist():
                content_xml = zf.read("content.xml")
                root = ET.fromstring(content_xml)
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        text_parts.append(elem.text.strip())
                    if elem.tail and elem.tail.strip():
                        text_parts.append(elem.tail.strip())
            else:
                notes.append("В ODT-файле не найден content.xml")

            for name in sorted(zf.namelist()):
                if name.startswith("Pictures/") and not name.startswith("Thumbnails/"):
                    ext = name.split(".")[-1].lower()
                    if ext not in IMAGE_EXTENSIONS:
                        notes.append(f"Пропущено изображение ODT неподдерживаемого формата: {name}")
                        continue
                    try:
                        raw = zf.read(name)
                        resized = _resize_image(raw)
                        images.append(("image/png", _image_to_base64(resized)))
                    except Exception as exc:
                        notes.append(f"Не удалось обработать изображение {name}: {exc}")
    except Exception as exc:
        notes.append(f"Не удалось открыть ODT как zip: {exc}")

    return ExtractedContent(text="\n".join(text_parts), images_base64=images, notes=notes)


def _extract_docx(content: bytes) -> ExtractedContent:
    if _is_odt(content):
        return _extract_odt(content)

    text_parts: List[str] = []
    images: List[Tuple[str, str]] = []
    notes: List[str] = []

    try:
        document = Document(io.BytesIO(content))
        for para in document.paragraphs:
            if para.text:
                text_parts.append(para.text)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_parts.append(cell.text)
    except Exception as exc:
        notes.append(f"Не удалось извлечь текст docx: {exc}")

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            for name in sorted(zf.namelist()):
                if name.startswith("word/media/"):
                    ext = name.split(".")[-1].lower()
                    if ext not in IMAGE_EXTENSIONS:
                        notes.append(f"Пропущено изображение docx неподдерживаемого формата: {name}")
                        continue
                    try:
                        raw = zf.read(name)
                        resized = _resize_image(raw)
                        images.append(("image/png", _image_to_base64(resized)))
                    except Exception as exc:
                        notes.append(f"Не удалось обработать изображение {name}: {exc}")
    except Exception as exc:
        notes.append(f"Не удалось открыть docx как zip: {exc}")

    return ExtractedContent(text="\n".join(text_parts), images_base64=images, notes=notes)


def _extract_pdf(content: bytes) -> ExtractedContent:
    text_parts: List[str] = []
    images: List[Tuple[str, str]] = []
    notes: List[str] = []

    try:
        reader = PdfReader(io.BytesIO(content))
        for page in reader.pages:
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            except Exception as exc:
                notes.append(f"Не удалось извлечь текст страницы: {exc}")

            try:
                for image in page.images:
                    try:
                        resized = _resize_image(image.data)
                        images.append(("image/png", _image_to_base64(resized)))
                    except Exception as exc:
                        notes.append(f"Не удалось обработать изображение pdf: {exc}")
            except Exception:
                pass
    except Exception as exc:
        notes.append(f"Не удалось прочитать PDF: {exc}")

    return ExtractedContent(text="\n".join(text_parts), images_base64=images, notes=notes)


def _extract_xlsx(content: bytes) -> ExtractedContent:
    text_parts: List[str] = []
    notes: List[str] = []
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join(str(cell) for cell in row if cell is not None)
                if row_text.strip():
                    text_parts.append(row_text)
    except Exception as exc:
        notes.append(f"Не удалось прочитать xlsx: {exc}")
    return ExtractedContent(text="\n".join(text_parts), images_base64=[], notes=notes)


def extract_file(key: str, original_filename: str) -> ExtractedContent:
    data = s3.get_object_bytes(key)
    ext = original_filename.split(".")[-1].lower() if "." in original_filename else ""

    if ext in ("docx", "doc", "odt"):
        return _extract_docx(data)
    if ext == "pdf":
        return _extract_pdf(data)
    if ext in ("xlsx", "xls"):
        return _extract_xlsx(data)
    if ext in IMAGE_EXTENSIONS:
        try:
            resized = _resize_image(data)
            return ExtractedContent(
                text="",
                images_base64=[("image/png", _image_to_base64(resized))],
                notes=[],
            )
        except Exception as exc:
            return ExtractedContent(text="", images_base64=[], notes=[f"Не удалось обработать изображение: {exc}"])
    if ext in TEXT_EXTENSIONS or not ext:
        return ExtractedContent(text=_read_text_bytes(data), images_base64=[], notes=[])

    return ExtractedContent(
        text="",
        images_base64=[],
        notes=[f"Формат файла .{ext} не поддерживается для извлечения содержимого."],
    )
